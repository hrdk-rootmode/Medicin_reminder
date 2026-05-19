#!/usr/bin/env python3
import sys
import os
import sqlite3
from pathlib import Path

try:
    from openpyxl import load_workbook
except Exception as e:
    print('openpyxl not installed:', e)
    print('You can install it with: python -m pip install openpyxl')
    sys.exit(2)

ROOT = Path(__file__).resolve().parents[1]
INPUT_XLSX = ROOT / 'india_medicines.xlsx'
OUTPUT_DB = ROOT / 'app' / 'src' / 'main' / 'assets' / 'india_medicines.db'

if not INPUT_XLSX.exists():
    print('Input file not found at', INPUT_XLSX)
    sys.exit(1)

OUTPUT_DB.parent.mkdir(parents=True, exist_ok=True)

wb = load_workbook(INPUT_XLSX, read_only=True)
ws = wb.active
rows = ws.iter_rows(values_only=True)

headers = []
try:
    headers = [str(c).strip() for c in next(rows)]
except StopIteration:
    print('Empty workbook')
    sys.exit(1)

print('Detected headers:', headers)

# heuristics to find columns
def find_col(names):
    names = [n.lower() for n in names]
    for idx, h in enumerate(names):
        if 'brand' in h or 'medicine' in h or 'name' in h or 'product' in h:
            return idx
    return None

brand_idx = find_col(headers)
# try generic
generic_idx = None
for idx, h in enumerate(headers):
    if 'generic' in h or 'generic_name' in h or 'pron' in h:
        generic_idx = idx
        break
# try strength
strength_idx = None
for idx, h in enumerate(headers):
    if 'strength' in h or 'dose' in h or 'mg' in h:
        strength_idx = idx
        break

if brand_idx is None:
    brand_idx = 0
if generic_idx is None:
    generic_idx = 1 if len(headers) > 1 else brand_idx
if strength_idx is None:
    strength_idx = 2 if len(headers) > 2 else None

print('Using cols - brand:', brand_idx, 'generic:', generic_idx, 'strength:', strength_idx)

entries = []
for r in rows:
    if not r:
        continue
    def val(i):
        try:
            return (str(r[i]).strip()) if r[i] is not None else ''
        except Exception:
            return ''
    brand = val(brand_idx)
    generic = val(generic_idx)
    strength = val(strength_idx) if strength_idx is not None else ''
    name = brand or generic
    if not name:
        continue
    entries.append((name, brand, generic, strength))

print('Parsed', len(entries), 'rows')

# normalization
import re

def normalize(s):
    if not s:
        return ''
    s = s.lower()
    s = re.sub(r"[^a-z0-9 ]+", ' ', s)
    s = re.sub(r"\s+", ' ', s).strip()
    return s

# create sqlite with fts5
if OUTPUT_DB.exists():
    try:
        OUTPUT_DB.unlink()
    except Exception as e:
        print('Failed to remove existing DB:', e)

conn = sqlite3.connect(str(OUTPUT_DB))
cur = conn.cursor()
# enable fts5
cur.execute("CREATE VIRTUAL TABLE medicines USING fts5(name, brand, generic, strength, normalized, tokenize='unicode61')")

insert_sql = 'INSERT INTO medicines(name, brand, generic, strength, normalized) VALUES (?, ?, ?, ?, ?)'
count = 0
for name, brand, generic, strength in entries:
    norm = normalize(' '.join([name, brand, generic, strength]))
    cur.execute(insert_sql, (name, brand, generic, strength, norm))
    count += 1

conn.commit()
conn.close()
print('Wrote', count, 'entries to', OUTPUT_DB)
print('Done')
